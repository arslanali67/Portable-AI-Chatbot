"""File ingestion + dedup tests — txt/md/pdf/docx, limits, security."""

import io
import uuid

import pytest
from fastapi.testclient import TestClient

from app.main import app

pytestmark = pytest.mark.identity

client = TestClient(app)

PASSWORD = "strong-password-123"
_RUN = uuid.uuid4().hex[:8]


def _email(name: str) -> str:
    return f"{name}-{_RUN}@example.com"


def _slug(name: str) -> str:
    return f"{name}-{_RUN}"


def _register(email: str, full_name: str = "File Tester"):
    r = client.post(
        "/api/v1/auth/register",
        json={"email": email, "password": PASSWORD, "full_name": full_name},
    )
    assert r.status_code == 201, r.text
    return r.json()


def _login(email: str) -> str:
    r = client.post("/api/v1/auth/login", data={"username": email, "password": PASSWORD})
    assert r.status_code == 200, r.text
    return r.json()["access_token"]


def _auth(token: str) -> dict[str, str]:
    return {"Authorization": f"Bearer {token}"}


def _setup() -> tuple[str, int, int]:
    email = _email(f"owner{uuid.uuid4().hex[:6]}")
    token = _login(_register(email)["email"])
    r = client.post(
        "/api/v1/organizations",
        json={"name": "Org", "slug": _slug(f"org{uuid.uuid4().hex[:6]}")},
        headers=_auth(token),
    )
    org_id = r.json()["id"]
    r = client.post(
        f"/api/v1/organizations/{org_id}/chatbots",
        json={"name": "Bot", "slug": _slug(f"bot{uuid.uuid4().hex[:6]}")},
        headers=_auth(token),
    )
    return token, org_id, r.json()["id"]


def _upload(token: str, org_id: int, bot_id: int, filename: str, content: bytes, title=None):
    files = {"file": (filename, content)}
    data = {"title": title} if title else None
    return client.post(
        f"/api/v1/organizations/{org_id}/chatbots/{bot_id}/knowledge/documents/file",
        files=files,
        data=data,
        headers=_auth(token),
    )


def _text_ingest(token: str, org_id: int, bot_id: int, content: str):
    return client.post(
        f"/api/v1/organizations/{org_id}/chatbots/{bot_id}/knowledge/documents",
        json={"name": "Doc", "content": content, "source_type": "text"},
        headers=_auth(token),
    )


def _simple_pdf_bytes() -> bytes:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("helvetica", size=12)
    pdf.cell(text="Hello from PDF")
    return bytes(pdf.output())


def _simple_docx_bytes() -> bytes:
    import docx

    buffer = io.BytesIO()
    document = docx.Document()
    document.add_paragraph("Hello from DOCX")
    document.save(buffer)
    return buffer.getvalue()


# --- File ingestion ---


def test_upload_txt() -> None:
    token, org_id, bot_id = _setup()
    r = _upload(token, org_id, bot_id, "notes.txt", b"Plain text file content.")
    assert r.status_code == 201
    body = r.json()
    assert body["source_type"] == "file"
    assert body["original_filename"] == "notes.txt"
    assert body["status"] == "ready"
    assert body["chunk_count"] >= 1


def test_upload_markdown() -> None:
    token, org_id, bot_id = _setup()
    r = _upload(token, org_id, bot_id, "readme.md", b"# Title\n\nSome markdown body.")
    assert r.status_code == 201
    assert r.json()["status"] == "ready"


def test_upload_pdf() -> None:
    token, org_id, bot_id = _setup()
    r = _upload(token, org_id, bot_id, "doc.pdf", _simple_pdf_bytes())
    assert r.status_code == 201
    assert r.json()["status"] == "ready"


def test_upload_docx() -> None:
    token, org_id, bot_id = _setup()
    r = _upload(token, org_id, bot_id, "doc.docx", _simple_docx_bytes())
    assert r.status_code == 201
    assert r.json()["status"] == "ready"


def test_unsupported_extension_rejected() -> None:
    token, org_id, bot_id = _setup()
    r = _upload(token, org_id, bot_id, "evil.xlsx", b"binary")
    assert r.status_code == 422


def test_empty_file_rejected() -> None:
    token, org_id, bot_id = _setup()
    r = _upload(token, org_id, bot_id, "empty.txt", b"")
    assert r.status_code == 422


def test_empty_extracted_text_rejected() -> None:
    token, org_id, bot_id = _setup()
    r = _upload(token, org_id, bot_id, "blank.txt", b"   \n  ")
    assert r.status_code == 422


def test_oversized_file_rejected() -> None:
    token, org_id, bot_id = _setup()
    big = b"x" * (10 * 1024 * 1024 + 1)
    r = _upload(token, org_id, bot_id, "big.txt", big)
    assert r.status_code == 413


def test_malformed_pdf_safe_error() -> None:
    token, org_id, bot_id = _setup()
    r = _upload(token, org_id, bot_id, "bad.pdf", b"not a pdf at all")
    assert r.status_code == 422


def test_malformed_docx_safe_error() -> None:
    token, org_id, bot_id = _setup()
    r = _upload(token, org_id, bot_id, "bad.docx", b"not a docx")
    assert r.status_code == 422


def test_path_traversal_safe() -> None:
    token, org_id, bot_id = _setup()
    r = _upload(token, org_id, bot_id, "../../etc/passwd.txt", b"content")
    assert r.status_code in (201, 422)
    if r.status_code == 201:
        assert r.json()["original_filename"] == "passwd.txt"


# --- Deduplication ---


def test_duplicate_content_same_chatbot_409() -> None:
    token, org_id, bot_id = _setup()
    assert _text_ingest(token, org_id, bot_id, "duplicate me please").status_code == 201
    assert _text_ingest(token, org_id, bot_id, "duplicate me please").status_code == 409


def test_same_content_different_chatbot_allowed() -> None:
    token, org_id, bot_a = _setup()
    r = client.post(
        f"/api/v1/organizations/{org_id}/chatbots",
        json={"name": "Bot2", "slug": _slug(f"bot2{uuid.uuid4().hex[:6]}")},
        headers=_auth(token),
    )
    bot_b = r.json()["id"]
    assert _text_ingest(token, org_id, bot_a, "shared content").status_code == 201
    assert _text_ingest(token, org_id, bot_b, "shared content").status_code == 201


def test_same_content_different_org_allowed() -> None:
    token_a, org_a, bot_a = _setup()
    token_b, org_b, bot_b = _setup()
    assert _text_ingest(token_a, org_a, bot_a, "cross org content").status_code == 201
    assert _text_ingest(token_b, org_b, bot_b, "cross org content").status_code == 201


def test_client_hash_injection_ignored() -> None:
    token, org_id, bot_id = _setup()
    r = client.post(
        f"/api/v1/organizations/{org_id}/chatbots/{bot_id}/knowledge/documents",
        json={
            "name": "Doc",
            "content": "hash injection test",
            "source_type": "text",
            "content_hash": "deadbeef",
        },
        headers=_auth(token),
    )
    assert r.status_code == 422  # extra="forbid"


# --- Security ---


def test_cross_org_upload_denied() -> None:
    token_a, org_a, bot_a = _setup()
    token_b, org_b, bot_b = _setup()
    r = client.post(
        f"/api/v1/organizations/{org_b}/chatbots/{bot_b}/knowledge/documents/file",
        files={"file": ("x.txt", b"content")},
        headers=_auth(token_a),
    )
    assert r.status_code == 403


def test_client_cannot_inject_vector_or_status() -> None:
    token, org_id, bot_id = _setup()
    r = client.post(
        f"/api/v1/organizations/{org_id}/chatbots/{bot_id}/knowledge/documents",
        json={"name": "D", "content": "x", "source_type": "text", "status": "ready"},
        headers=_auth(token),
    )
    assert r.status_code == 422
